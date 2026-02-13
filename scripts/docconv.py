#!/usr/bin/env python3
"""
Утилита конвертации документов для DS Team Management.

Поддерживает:
  read-docx  FILE           → извлечь текст из .docx в stdout (markdown)
  read-xlsx  FILE [SHEET]   → извлечь данные из .xlsx в stdout (markdown-таблица)
  md-to-docx SRC [DST]      → конвертировать .md → .docx (pandoc)
  md-to-pdf  SRC [DST]      → конвертировать .md → .pdf  (pandoc, если pdflatex есть)
"""
import sys
import os
import json


def read_docx(path: str) -> str:
    """Извлекает текст из .docx и возвращает как markdown."""
    from docx import Document

    doc = Document(path)
    lines = []

    for para in doc.paragraphs:
        style = para.style.name.lower()
        text = para.text.strip()
        if not text:
            lines.append("")
            continue

        if "heading 1" in style:
            lines.append(f"# {text}")
        elif "heading 2" in style:
            lines.append(f"## {text}")
        elif "heading 3" in style:
            lines.append(f"### {text}")
        elif "list" in style or text.startswith(("-", "•", "–")):
            lines.append(f"- {text.lstrip('-•– ')}")
        else:
            lines.append(text)

    for table in doc.tables:
        lines.append("")
        for i, row in enumerate(table.rows):
            cells = [cell.text.strip().replace("|", "\\|") for cell in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
            if i == 0:
                lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
        lines.append("")

    return "\n".join(lines)


def read_xlsx(path: str, sheet_name: str | None = None) -> str:
    """Извлекает данные из .xlsx и возвращает как markdown-таблицы."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = [sheet_name] if sheet_name else wb.sheetnames
    output_parts = []

    for sname in sheets:
        if sname not in wb.sheetnames:
            output_parts.append(f"## Лист '{sname}' не найден\n")
            continue

        ws = wb[sname]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            output_parts.append(f"## {sname}\n\n_Пустой лист_\n")
            continue

        output_parts.append(f"## {sname}\n")

        header = rows[0]
        col_names = [str(c) if c is not None else "" for c in header]
        output_parts.append("| " + " | ".join(col_names) + " |")
        output_parts.append("| " + " | ".join(["---"] * len(col_names)) + " |")

        for row in rows[1:]:
            cells = [str(c) if c is not None else "" for c in row]
            # Pad if row has fewer cells
            while len(cells) < len(col_names):
                cells.append("")
            output_parts.append("| " + " | ".join(cells[:len(col_names)]) + " |")

        output_parts.append("")

    wb.close()
    return "\n".join(output_parts)


def md_to_docx(src: str, dst: str | None = None) -> str:
    """Конвертирует .md → .docx через pandoc."""
    import pypandoc

    if dst is None:
        dst = os.path.splitext(src)[0] + ".docx"

    pypandoc.convert_file(
        src,
        "docx",
        outputfile=dst,
        extra_args=["--standalone"],
    )
    return dst


def md_to_pdf(src: str, dst: str | None = None) -> str:
    """Конвертирует .md → .pdf через pandoc (требует pdflatex)."""
    import pypandoc

    if dst is None:
        dst = os.path.splitext(src)[0] + ".pdf"

    pypandoc.convert_file(
        src,
        "pdf",
        outputfile=dst,
        extra_args=["--standalone", "--pdf-engine=pdflatex"],
    )
    return dst


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)

    cmd = sys.argv[1]
    path = sys.argv[2]

    if not os.path.isfile(path):
        print(f"Файл не найден: {path}", file=sys.stderr)
        sys.exit(1)

    if cmd == "read-docx":
        print(read_docx(path))

    elif cmd == "read-xlsx":
        sheet = sys.argv[3] if len(sys.argv) > 3 else None
        print(read_xlsx(path, sheet))

    elif cmd == "md-to-docx":
        dst = sys.argv[3] if len(sys.argv) > 3 else None
        result = md_to_docx(path, dst)
        print(f"Конвертировано: {path} → {result}")

    elif cmd == "md-to-pdf":
        dst = sys.argv[3] if len(sys.argv) > 3 else None
        result = md_to_pdf(path, dst)
        print(f"Конвертировано: {path} → {result}")

    else:
        print(f"Неизвестная команда: {cmd}", file=sys.stderr)
        print(__doc__)
        sys.exit(1)


if __name__ == "__main__":
    main()
